import os
import sys
import datetime
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_callout(doc, text_paragraphs, title="NOTE", bg_hex="F0F4F8", border_hex="1D4ED8"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_before = Pt(2)
    p0.paragraph_format.space_after = Pt(4)
    run_title = p0.add_run(f"[{title}] ")
    run_title.bold = True
    run_title.font.name = "Arial"
    run_title.font.size = Pt(10)
    run_title.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8) if border_hex == "1D4ED8" else RGBColor(0xB9, 0x1C, 0x1C)
    
    if isinstance(text_paragraphs, str):
        text_paragraphs = [text_paragraphs]
        
    for i, t in enumerate(text_paragraphs):
        if i == 0:
            r = p0.add_run(t)
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        else:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(t)
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def style_table(table, col_widths, header_bg="1E3A8A", alt_bg="F9FAFB"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Header row
    hdr_cells = table.rows[0].cells
    for i, cell in enumerate(hdr_cells):
        cell.width = col_widths[i]
        set_cell_background(cell, header_bg)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.bold = True
                r.font.name = "Arial"
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                
    # Data rows
    for r_idx in range(1, len(table.rows)):
        row = table.rows[r_idx]
        bg = alt_bg if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(9.0)
                    r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

print("Helper functions defined successfully.")
